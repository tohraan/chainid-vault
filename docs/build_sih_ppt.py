from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = "/Users/tohraan/Downloads/SIH2026-IDEA-Presentation-Format.pptx"
OUT = "/Users/tohraan/Downloads/build/docs/SIH2026_ChainIDVault_IdeaPPT.pptx"

TNR = "Times New Roman"
INK  = RGBColor(0x00,0x00,0x00)
NAVY = RGBColor(0x1F,0x30,0x64)
MUTE = RGBColor(0x5A,0x57,0x51)
LINE = RGBColor(0xBF,0xBF,0xBF)
SOFT = RGBColor(0xF2,0xF2,0xF2)
RED  = RGBColor(0xA6,0x1B,0x1B)
REDBG= RGBColor(0xFB,0xE9,0xE9)
GRN  = RGBColor(0x1E,0x6B,0x3A)
GRNBG= RGBColor(0xE8,0xF3,0xEC)
AMB  = RGBColor(0xB4,0x53,0x09)
WHITE= RGBColor(0xFF,0xFF,0xFF)

prs = Presentation(SRC)

def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buChar","a:buAutoNum","a:buNone"):
        for el in pPr.findall(qn(tag)): pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))
    pPr.set("indent","0"); pPr.set("marL","0")

def clear(tf):
    tf.clear(); return tf

def add(tf, text, size, bold=False, first=False, before=0, color=INK, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    if before: p.space_before = Pt(before)
    p.space_after = Pt(2)
    if line: p.line_spacing = line
    r = p.add_run(); r.text = text
    r.font.name = TNR; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    no_bullet(p)
    return p

def shape_of(slide, name):
    for sh in slide.shapes:
        if sh.name == name: return sh

def txt(s, x, y, w, h, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT,
        line=None, anchor=MSO_ANCHOR.TOP, font=TNR):
    tb = s.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(1)
        if line: p.line_spacing = line
        r = p.add_run(); r.text = ln
        r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def box(s, x, y, w, h, fill=WHITE, border=LINE, rounded=False):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                            In(x), In(y), In(w), In(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border is None: sh.line.fill.background()
    else: sh.line.color.rgb = border; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh

def arrow(s, x, y, w=0.35, h=0.16, color=NAVY, down=False):
    sh = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW if down else MSO_SHAPE.RIGHT_ARROW,
                            In(x), In(y), In(w), In(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background(); sh.shadow.inherit=False
    return sh

def panel_title(s, x, y, w, t):
    txt(s, x, y, w, 0.22, t.upper(), 10, True, NAVY)

def oval(s, team="[TEAM NAME]"):
    for sh in s.shapes:
        if sh.name.startswith("Oval"):
            tf = clear(sh.text_frame); add(tf, team, 10, True, first=True, color=NAVY)

# geometry: left column for text, right column for the diagram
LX, LW = 0.55, 6.05
RX, RW = 6.95, 5.75
TOP = 1.15

# ═══════════════════════════════════ SLIDE 1 — title
s = prs.slides[0]
tf = clear(shape_of(s,"TextBox 9").text_frame)
rows = [("Problem Statement ID – ","SIH26125"),
        ("Problem Statement Title – ","Blockchain-Based Secure Platform for Identity, Access Control, and Digital Asset Management"),
        ("Theme – ","Blockchain & Cybersecurity"),
        ("PS Category – ","Software"),
        ("Team ID – ","[ENTER TEAM ID]"),
        ("Team Name (Registered on portal) – ","[ENTER TEAM NAME]")]
for i,(a,b) in enumerate(rows):
    p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
    p.space_after = Pt(9); no_bullet(p)
    r1=p.add_run(); r1.text=a; r1.font.name=TNR; r1.font.size=Pt(16); r1.font.bold=True; r1.font.color.rgb=INK
    r2=p.add_run(); r2.text=b; r2.font.name=TNR; r2.font.size=Pt(16); r2.font.color.rgb=INK
tf = clear(shape_of(s,"Subtitle 3").text_frame)
add(tf,"ChainID Vault",28,True,first=True,color=NAVY)
add(tf,"Tamper-proof custody of identity, access and assets",15)

# ═══════════════════════════════════ SLIDE 2 — IDEA
s = prs.slides[1]; oval(s)
tb = shape_of(s,"TextBox 8"); tb.left,tb.top,tb.width,tb.height = In(LX),In(TOP),In(LW),In(5.15)
tf = clear(tb.text_frame)
add(tf,"ChainID Vault — identity, access and asset custody enforced inside smart contracts, not inside an editable database.",13,True,first=True,color=NAVY,line=1.15)
add(tf,"Proposed Solution (Describe your Idea/Solution/Prototype)",13,True,before=8)
for b in ["Every identity, role grant and custody record is written on a permissioned blockchain.",
          "Identity = keypair + lifecycle status (Active / Suspended / Revoked) + hash of an off-chain record.",
          "Asset = ERC-721 token for one controlled item, bound to a verified identity.",
          "Sensitive documents stay off-chain; only a keccak256 fingerprint is anchored."]:
    add(tf,"•  "+b,11,line=1.1)
add(tf,"How it addresses the problem",13,True,before=7)
for b in ["No database row exists for a role, an owner or a log entry, so none can be edited.",
          "A revoked identity cannot receive an asset by any route — one enforcement chokepoint.",
          "The audit trail is the chain's append-only event log; it cannot be selectively deleted.",
          "Any third party verifies an asset with no account and no trust in the operator."]:
    add(tf,"•  "+b,11,line=1.1)
add(tf,"Innovation and uniqueness of the solution",13,True,before=7)
for b in ["Single-chokepoint enforcement closes the ERC-721 transfer bypass most implementations leave open.",
          "EIP-712 proof of control: single-use, time-bound, domain-separated — replay-proof.",
          "Even the administrator who created an identity cannot override the rules."]:
    add(tf,"•  "+b,11,line=1.1)

panel_title(s,RX,TOP,RW,"The three governed entities")
ent=[("IDENTITY","Recognised keypair\nActive / Suspended / Revoked",NAVY),
     ("ASSET","ERC-721 token for one\ncontrolled physical item",GRN),
     ("ROLE","Admin · Manager\nAuditor · User",AMB)]
yy=TOP+0.3
for n,d,c in ent:
    box(s,RX,yy,RW,0.82,fill=WHITE,border=c)
    box(s,RX,yy,0.09,0.82,fill=c,border=None)
    txt(s,RX+0.25,yy+0.11,RW-0.4,0.22,n,12,True,c)
    txt(s,RX+0.25,yy+0.36,RW-0.4,0.42,d,10.5,False,MUTE,line=1.12)
    yy+=0.95
arrow(s,RX+RW/2-0.1,TOP+0.86,0.2,0.09,NAVY,down=True)
arrow(s,RX+RW/2-0.1,TOP+1.81,0.2,0.09,NAVY,down=True)
box(s,RX,yy+0.05,RW,1.5,fill=SOFT,border=LINE)
txt(s,RX+0.22,yy+0.22,RW-0.44,0.25,"THE INVARIANT",10,True,NAVY)
txt(s,RX+0.22,yy+0.5,RW-0.44,0.9,
    "An asset may only ever be held by an\nactive identity — enforced in _update, the\none function every mint, transfer and burn\npasses through.",11,False,INK,line=1.18)
box(s,RX,yy+1.68,RW,0.72,fill=GRNBG,border=GRN)
txt(s,RX+0.22,yy+1.84,RW-0.44,0.45,"Working prototype · 45 automated contract tests passing\n6 screens · full custody workflow demonstrated live",10.5,True,GRN,line=1.15)

# ═══════════════════════════════════ SLIDE 3 — TECHNICAL
s = prs.slides[2]; oval(s)
tb = shape_of(s,"TextBox 8"); tb.left,tb.top,tb.width,tb.height = In(LX),In(TOP),In(LW),In(5.15)
tf = clear(tb.text_frame)
add(tf,"Technologies to be used (e.g. programming languages, frameworks, hardware)",13,True,first=True)
for b in ["Contracts: Solidity 0.8.24, OpenZeppelin 5.6 (AccessControl, ERC-721 Enumerable), Hardhat.",
          "Frontend: React 18, Vite, Tailwind, ethers.js v6 — the browser calls contracts directly.",
          "Cryptography: EIP-712 typed-data signatures; keccak256 document anchoring.",
          "No backend server and no database — deliberately, see the architecture opposite.",
          "Deployment target: permissioned EVM chain (Hyperledger Besu / Quorum)."]:
    add(tf,"•  "+b,11,line=1.1)
add(tf,"Methodology and process for implementation (Flow Charts / Images / working prototype)",13,True,before=8)
for b in ["Two contracts: IdentityRegistry governs lifecycle and proof of control; AssetNFT governs custody and calls the registry before every movement.",
          "The full contractor custody pipeline runs live, end to end, in the prototype.",
          "45 tests including negative cases: replay, impersonation, expiry, tampering."]:
    add(tf,"•  "+b,11,line=1.1)
add(tf,"Verified refusals demonstrated on-chain",12,True,before=8,color=RED)
for b in ["Non-admin cannot issue · non-admin cannot suspend an identity",
          "Suspended identity cannot receive · tampered document fails verification"]:
    add(tf,"•  "+b,11,line=1.1,color=RED)

panel_title(s,RX,TOP,RW,"System architecture")
box(s,RX,TOP+0.3,RW,0.75,fill=WHITE,border=LINE)
txt(s,RX+0.2,TOP+0.42,RW-0.4,0.22,"BROWSER — React + ethers.js",11,True,INK)
txt(s,RX+0.2,TOP+0.68,RW-0.4,0.3,"Six screens. Holds no authority.",10,False,MUTE)
arrow(s,RX+RW/2-0.1,TOP+1.1,0.2,0.16,NAVY,down=True)
txt(s,RX+RW/2+0.16,TOP+1.1,1.6,0.2,"JSON-RPC",9,True,MUTE)
box(s,RX,TOP+1.35,RW,1.25,fill=NAVY,border=None)
txt(s,RX+0.2,TOP+1.48,RW-0.4,0.22,"PERMISSIONED CHAIN",10,True,WHITE)
txt(s,RX+0.2,TOP+1.75,RW-0.4,0.22,"IdentityRegistry.sol",11.5,True,WHITE)
txt(s,RX+0.2,TOP+2.0,RW-0.4,0.22,"AssetNFT.sol",11.5,True,WHITE)
txt(s,RX+0.2,TOP+2.28,RW-0.4,0.22,"All authority lives here",9.5,False,RGBColor(0xC8,0xCE,0xDC))
arrow(s,RX+RW/2-0.1,TOP+2.65,0.2,0.16,NAVY,down=True)
box(s,RX,TOP+2.9,RW,0.72,fill=SOFT,border=LINE)
txt(s,RX+0.2,TOP+3.02,RW-0.4,0.22,"ANY VERIFIER — no account required",11,True,INK)
txt(s,RX+0.2,TOP+3.28,RW-0.4,0.3,"Reads the same record without trusting us.",10,False,MUTE)
box(s,RX,TOP+3.78,RW,1.35,fill=REDBG,border=RED)
txt(s,RX+0.2,TOP+3.92,RW-0.4,0.22,"DELIBERATELY ABSENT",10,True,RED)
for i,(a,b) in enumerate([("No backend server","nothing between user and chain"),
                          ("No database","no editable row for role, owner or log"),
                          ("No public network","permissioned chain, no gas, no exposure")]):
    txt(s,RX+0.2,TOP+4.18+i*0.30,2.0,0.22,a,10.5,True,INK)
    txt(s,RX+2.15,TOP+4.19+i*0.30,RW-2.35,0.22,b,9.5,False,MUTE)

prs.save(OUT); print("slides 1-3 done")

# ═══════════════════════════════════ SLIDE 4 — FEASIBILITY
s = prs.slides[3]; oval(s)
tb = shape_of(s,"TextBox 8"); tb.left,tb.top,tb.width,tb.height = In(LX),In(TOP),In(LW),In(5.15)
tf = clear(tb.text_frame)
add(tf,"Analysis of the feasibility of the idea",13,True,first=True)
for b in ["Working prototype exists: two contracts, six screens, 45 passing tests.",
          "Built on audited OpenZeppelin libraries — no novel cryptography.",
          "Runs on a permissioned chain: no public network, no gas cost, no external dependency.",
          "Adoption is incremental — can run alongside an existing stores register before replacing it."]:
    add(tf,"•  "+b,11,line=1.1)
add(tf,"Potential challenges and risks",13,True,before=8)
for b in ["Administrator key compromise — one key currently holds full authority.",
          "On-chain read cost grows with registry size.",
          "The chain proves the integrity of the record, not the identity of the physical object.",
          "Key loss would otherwise strand an identity and the assets it holds."]:
    add(tf,"•  "+b,11,line=1.1)
add(tf,"Strategies for overcoming these challenges",13,True,before=8)
for b in ["Multi-signature control with a timelock for the administrative role.",
          "Paginated on-chain reads (implemented) plus an off-chain indexer at scale.",
          "Tamper-evident physical tagging bound to the asset identifier.",
          "On-chain key rotation preserving identity continuity (implemented)."]:
    add(tf,"•  "+b,11,line=1.1)

panel_title(s,RX,TOP,RW,"The documented case this is built for")
box(s,RX,TOP+0.28,RW,1.62,fill=REDBG,border=RED)
txt(s,RX+0.2,TOP+0.42,RW-0.4,0.25,"Ordnance Factory Board (now seven defence PSUs)",12,True,RED)
txt(s,RX+0.2,TOP+0.72,RW-0.4,1.05,
    "CAG Audit Report No. 10 of 2024 found stores-in-hand of\n₹6,172 crore at 31 March 2021, of which ₹1,237 crore was\nnon-active — non-moving, slow-moving, surplus, scrap or\nobsolete. Holding rose from 214 to 344 days of consumption.",10.5,False,INK,line=1.2)
box(s,RX,TOP+2.0,RW,0.82,fill=REDBG,border=RED)
txt(s,RX+0.2,TOP+2.13,RW-0.4,0.25,"And it persists after corporatisation",11.5,True,RED)
txt(s,RX+0.2,TOP+2.4,RW-0.4,0.38,
    "CAG Report No. 34: 21% of stores-in-hand — ₹330 crore across\nfour sampled factories — still non-active at 31 March 2023.",10.5,False,INK,line=1.18)
box(s,RX,TOP+2.94,RW,1.05,fill=SOFT,border=LINE)
txt(s,RX+0.2,TOP+3.07,RW-0.4,0.25,"WHY THIS IS OUR PROBLEM STATEMENT",10,True,NAVY)
txt(s,RX+0.2,TOP+3.34,RW-0.4,0.58,
    "Stock nobody can reconcile is a custody-record failure, not a\nwarehouse failure. The register and the physical reality diverged,\nand no party could independently prove which was correct.",10.5,False,INK,line=1.2)
box(s,RX,TOP+4.11,RW,1.02,fill=GRNBG,border=GRN)
txt(s,RX+0.2,TOP+4.24,RW-0.4,0.25,"THE ACCESS SIDE OF THE SAME PROBLEM",10,True,GRN)
txt(s,RX+0.2,TOP+4.51,RW-0.4,0.55,
    "Only 48% of organisations know whether former staff still hold\naccess; 20% suffered a breach as a result (USENIX Security, 2025).\nStolen credentials are the top initial vector — 22% (Verizon DBIR 2025).",10,False,INK,line=1.18)

# ═══════════════════════════════════ SLIDE 5 — IMPACT
s = prs.slides[4]; oval(s)
tb = shape_of(s,"TextBox 8"); tb.left,tb.top,tb.width,tb.height = In(LX),In(TOP),In(LW),In(5.15)
tf = clear(tb.text_frame)
add(tf,"Potential impact on the target audience",13,True,first=True)
for b in ["Security administrators: revocation is immediate and cannot be bypassed, including by them.",
          "Stores and custody officers: one verifiable record instead of a register and paperwork that disagree.",
          "Internal auditors: full history reconstructable without trusting the operator's own copy.",
          "External verifiers and receiving officers: confirm authenticity with no account and no login."]:
    add(tf,"•  "+b,11,line=1.1)
add(tf,"Benefits of the solution (social, economic, environmental, etc.)",13,True,before=8)
for b in ["Security: authorisation lives in contract code — a compromised database cannot grant privilege or rewrite history.",
          "Economic: directly targets the reconciliation failure behind ₹1,237 crore of non-active defence stores (CAG, 2024).",
          "Operational: enforced offboarding closes the gap that opens every time somebody leaves.",
          "Compliance and privacy: tamper-evident custody chain, while personal data stays off-chain and erasable.",
          "Strategic: a defensible custody trail for sensitive equipment across sites, contractors and vendors."]:
    add(tf,"•  "+b,11,line=1.1)

panel_title(s,RX,TOP,RW,"Before and after, on the same KPIs")
hdr_y = TOP+0.3
box(s,RX,hdr_y,RW,0.3,fill=NAVY,border=None)
txt(s,RX+0.12,hdr_y+0.06,2.05,0.2,"KPI",9.5,True,WHITE)
txt(s,RX+2.2,hdr_y+0.06,1.7,0.2,"TODAY",9.5,True,WHITE)
txt(s,RX+4.0,hdr_y+0.06,1.7,0.2,"WITH CHAINID VAULT",9.5,True,WHITE)
rows=[("Custody record","Editable row","Append-only ledger"),
      ("Reconciliation","Periodic, manual","Continuous — the ledger is the record"),
      ("Offboarding gap","Access removed,\nassets still held","Revoked identity cannot\nhold or receive"),
      ("Independent check","Not possible","Any party, no login"),
      ("Time to block a\nrevoked holder","Process-dependent","One transaction"),
      ("Audit evidence","Operator's own copy","Chain, verifiable by anyone")]
yy=hdr_y+0.3
for i,(k,a,b) in enumerate(rows):
    h = 0.46 if "\n" in (k+a+b) else 0.33
    box(s,RX,yy,RW,h,fill=WHITE if i%2==0 else SOFT,border=LINE)
    txt(s,RX+0.12,yy+0.06,2.0,h-0.1,k,9.5,True,INK,line=1.05)
    txt(s,RX+2.2,yy+0.06,1.75,h-0.1,a,9.5,False,RED,line=1.05)
    txt(s,RX+4.0,yy+0.06,1.65,h-0.1,b,9.5,False,GRN,line=1.05)
    yy+=h
box(s,RX,yy+0.12,RW,0.72,fill=SOFT,border=LINE)
txt(s,RX+0.15,yy+0.24,RW-0.3,0.5,
    "The right-hand column states properties the system enforces\nstructurally. They are design guarantees, not measured field results —\nthere is no deployment yet.",9.5,False,MUTE,line=1.18)
box(s,RX,yy+0.94,RW,0.62,fill=GRNBG,border=GRN)
txt(s,RX+0.15,yy+1.06,RW-0.3,0.42,
    "Evidenced today: 45 contract tests passing · 3 exploits found in our\nown code and closed · 4 distinct on-chain refusals demonstrated.",9.5,True,GRN,line=1.18)

# ═══════════════════════════════════ SLIDE 6 — REFERENCES
s = prs.slides[5]; oval(s)
tb = shape_of(s,"TextBox 8"); tb.left,tb.top,tb.width,tb.height = In(LX),In(TOP),In(LW),In(5.15)
tf = clear(tb.text_frame)
add(tf,"Details / Links of the reference and research work",13,True,first=True)
add(tf,"Problem evidence",12,True,before=6,color=NAVY)
for b in ["CAG Audit Report No. 10 of 2024 — Ordnance Factory Board inventory management. cag.gov.in",
          "CAG Audit Report No. 34 — Defence Public Sector Undertakings, stores-in-hand. cag.gov.in",
          "IBM Cost of a Data Breach Report 2025 — India. in.newsroom.ibm.com",
          "Verizon Data Breach Investigations Report 2025. verizon.com/business/resources/reports/dbir",
          "USENIX Security Symposium 2025 — former-employee access retention."]:
    add(tf,"•  "+b,10.5,line=1.1)
add(tf,"Technical standards",12,True,before=7,color=NAVY)
for b in ["ERC-721 Non-Fungible Token Standard — eips.ethereum.org/EIPS/eip-721",
          "EIP-712 Typed Structured Data Hashing and Signing — eips.ethereum.org/EIPS/eip-712",
          "OpenZeppelin Contracts v5 — docs.openzeppelin.com/contracts/5.x",
          "NIST SP 800-63-3 Digital Identity Guidelines — pages.nist.gov/800-63-3",
          "OWASP Smart Contract Top 10 — owasp.org",
          "Hyperledger Besu, permissioned EVM — besu.hyperledger.org",
          "W3C Decentralized Identifiers v1.0 — w3.org/TR/did-core (reviewed, not claimed)"]:
    add(tf,"•  "+b,10.5,line=1.1)
add(tf,"Prototype",12,True,before=7,color=NAVY)
add(tf,"•  Source, tests, security audit and threat model — github.com/tohraan/chainid-vault",10.5)

panel_title(s,RX,TOP,RW,"The number this is built on")
box(s,RX,TOP+0.3,RW,1.75,fill=NAVY,border=None)
txt(s,RX,TOP+0.5,RW,0.75,"₹1,237 crore",40,True,WHITE,PP_ALIGN.CENTER)
txt(s,RX+0.3,TOP+1.32,RW-0.6,0.6,
    "of non-active stores at the Ordnance Factory Board —\nstock the system could not reconcile.",11.5,False,RGBColor(0xC8,0xCE,0xDC),PP_ALIGN.CENTER,line=1.18)
txt(s,RX,TOP+2.12,RW,0.25,"CAG Audit Report No. 10 of 2024, Government of India",9.5,False,MUTE,PP_ALIGN.CENTER)
box(s,RX,TOP+2.55,RW,1.15,fill=WHITE,border=LINE)
txt(s,RX+0.2,TOP+2.68,RW-0.4,0.22,"AND THE SAME FAILURE, TWO YEARS LATER",10,True,RED)
txt(s,RX,TOP+2.95,RW,0.45,"21%  ·  ₹330 crore",22,True,RED,PP_ALIGN.CENTER)
txt(s,RX+0.2,TOP+3.4,RW-0.4,0.25,"of stores-in-hand still non-active, four sampled factories,\n31 March 2023 — CAG Report No. 34",9.5,False,MUTE,PP_ALIGN.CENTER,line=1.15)
box(s,RX,TOP+3.85,RW,1.28,fill=SOFT,border=LINE)
txt(s,RX+0.2,TOP+3.98,RW-0.4,0.22,"WHAT WE DO NOT CLAIM",10,True,NAVY)
txt(s,RX+0.2,TOP+4.24,RW-0.4,0.8,
    "We have not deployed and cannot claim field results. Every figure\nabove is the documented problem, not our measured effect. What we\ncan evidence is a working prototype, 45 passing tests, and three\nexploits found in our own code and closed.",10,False,INK,line=1.18)

# delete the instructions slide
xs = prs.slides._sldIdLst; sl = list(xs)
prs.part.drop_rel(sl[6].rId); xs.remove(sl[6])
prs.save(OUT)
print("saved", OUT, "slides:", len(Presentation(OUT).slides))
